from __future__ import annotations

import re
import zipfile
from io import BytesIO

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sam.suggest import ordered_ids
from sam.time_utils import clean_token, format_time, subtract_minutes, time_to_minutes
from sam.weekday import KEEP_COLUMNS, block_sort_key

PACKET_COLUMNS = [
    "Asn#",
    "Block",
    "Route",
    "Direction",
    "Start Time",
    "Start Location",
    "End Location",
    "End Time",
]
COLUMN_WIDTHS = [0.7, 0.7, 0.7, 0.95, 0.95, 2.45, 2.45, 0.95]


def _rows_for_label(frame: pd.DataFrame, label: str) -> list[dict[str, object]]:
    matched = frame.loc[frame["Asn#"].map(clean_token).eq(clean_token(label))]
    ids = ordered_ids(frame, [int(value) for value in matched["_id"].tolist()])
    rows: list[dict[str, object]] = []
    for row_id in ids:
        source = frame.loc[frame["_id"] == row_id].iloc[0]
        row = {column: source[column] for column in PACKET_COLUMNS}
        row["_id"] = int(row_id)
        rows.append(row)
    return rows


def assignment_labels(frame: pd.DataFrame) -> list[str]:
    labels = [clean_token(value) for value in frame["Asn#"].tolist() if clean_token(value)]
    unique = list(dict.fromkeys(labels))

    def sort_key(label: str) -> tuple:
        if label.isdigit():
            return (0, int(label), "")
        return (1, 0, label.casefold())

    return sorted(unique, key=sort_key)


def _safe_filename(text: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*#]+', "_", text)
    cleaned = re.sub(r"\s+", "_", cleaned).strip("._")
    return cleaned or "assignment"


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _display(header: str, value: object) -> str:
    if "Time" in header:
        return format_time(value)
    return clean_token(value)


def _configure_section(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.25)


def _add_header_line(header, label: str, value: str) -> None:
    paragraph = header.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    label_run = paragraph.add_run(f"{label}  ")
    label_run.bold = True
    label_run.font.name = "Calibri"
    label_run.font.size = Pt(11)
    label_run.font.color.rgb = RGBColor(15, 110, 86)
    value_run = paragraph.add_run(value or "—")
    value_run.font.name = "Calibri"
    value_run.font.size = Pt(11)


def _clear_header(header) -> None:
    for paragraph in list(header.paragraphs):
        element = paragraph._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)


def _append_packet(document: Document, label: str, rows: list[dict[str, object]], notes: list[dict[str, object]], *, new_page: bool) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE) if new_page else document.sections[0]
    _configure_section(section)
    first = rows[0]
    last = rows[-1]
    blocks = list(dict.fromkeys(clean_token(row["Block"]) for row in rows if clean_token(row["Block"])))
    header = section.header
    header.is_linked_to_previous = False
    _clear_header(header)
    _add_header_line(header, "Assignment", label)
    _add_header_line(header, "Report at", _display("Start Location", first.get("Start Location")))
    _add_header_line(header, "Report time", subtract_minutes(first.get("Start Time"), 15))
    _add_header_line(header, "Block", ", ".join(blocks))
    _add_header_line(
        header,
        "Shift",
        f"{format_time(first.get('Start Time'))} – {format_time(last.get('End Time'))}",
    )

    notes_after: dict[int, list[str]] = {}
    for note in notes:
        text = clean_token(note.get("text"))
        after_id = note.get("after_row_id")
        if text and after_id is not None:
            notes_after.setdefault(int(after_id), []).append(text)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(78)

    table = document.add_table(rows=1, cols=len(PACKET_COLUMNS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, name in enumerate(PACKET_COLUMNS):
        cell = table.rows[0].cells[index]
        cell.text = ""
        cell.width = Inches(COLUMN_WIDTHS[index])
        _shade(cell, "E7F3EF")
        run = cell.paragraphs[0].add_run(name)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(15, 110, 86)

    for row in rows:
        cells = table.add_row().cells
        for index, name in enumerate(PACKET_COLUMNS):
            cell = cells[index]
            cell.text = ""
            cell.width = Inches(COLUMN_WIDTHS[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            run = cell.paragraphs[0].add_run(_display(name, row.get(name)))
            run.font.name = "Calibri"
            run.font.size = Pt(9)
        for text in notes_after.get(int(row["_id"]), []):
            note_cells = table.add_row().cells
            note_cells[0].merge(note_cells[-1])
            note_cells[0].text = ""
            _shade(note_cells[0], "FFF4E5")
            run = note_cells[0].paragraphs[0].add_run(text)
            run.italic = True
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(146, 64, 14)


def _document_bytes(document: Document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _new_document() -> Document:
    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)
    return document


def packet_bytes(frame: pd.DataFrame, label: str, notes: dict[str, list[dict[str, object]]]) -> bytes:
    rows = _rows_for_label(frame, label)
    if not rows:
        raise ValueError(f"{label} has no trips yet.")
    document = _new_document()
    _append_packet(document, label, rows, notes.get(label, []), new_page=False)
    return _document_bytes(document)


def packets_zip(frame: pd.DataFrame, notes: dict[str, list[dict[str, object]]]) -> bytes:
    labels = assignment_labels(frame)
    if not labels:
        raise ValueError("No assignments to export yet.")
    combined = _new_document()
    archive = BytesIO()
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for index, label in enumerate(labels):
            rows = _rows_for_label(frame, label)
            label_notes = notes.get(label, [])
            single = _new_document()
            _append_packet(single, label, rows, label_notes, new_page=False)
            bundle.writestr(f"{_safe_filename(label)}.docx", _document_bytes(single))
            _append_packet(combined, label, rows, label_notes, new_page=index > 0)
        bundle.writestr("_ALL_Assignments.docx", _document_bytes(combined))
    return archive.getvalue()


def _excel_value(column: str, value: object):
    if "Time" in column:
        if value is None or pd.isna(value):
            return None
        if isinstance(value, pd.Timestamp):
            return value.to_pydatetime().replace(tzinfo=None)
        return value
    text = clean_token(value)
    return text or None


def workbook_bytes(frame: pd.DataFrame) -> bytes:
    book = Workbook()
    _write_sheet(book.active, "Weekday", frame, KEEP_COLUMNS, assigned_only=False)
    _write_sheet(book.create_sheet("Made assignments"), "Made assignments", frame, PACKET_COLUMNS, assigned_only=True)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _write_sheet(sheet, title: str, frame: pd.DataFrame, columns: list[str], assigned_only: bool) -> None:
    sheet.title = title
    header_fill = PatternFill("solid", fgColor="0F6E56")
    header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    source = frame.copy()
    if assigned_only:
        source = source.loc[source["Asn#"].map(clean_token).ne("")]
        source["_label_sort"] = source["Asn#"].map(
            lambda value: (0, int(clean_token(value)), "")
            if clean_token(value).isdigit()
            else (1, 0, clean_token(value).casefold())
        )
        source["_start_sort"] = source["Start Time"].map(lambda value: time_to_minutes(value) if time_to_minutes(value) is not None else 99999)
        source = source.sort_values(["_label_sort", "_start_sort", "_id"])
    else:
        source["_block_sort"] = source["Block"].map(lambda value: block_sort_key(clean_token(value)))
        source["_start_sort"] = source["Start Time"].map(lambda value: time_to_minutes(value) if time_to_minutes(value) is not None else 99999)
        source = source.sort_values(["_block_sort", "_start_sort", "_id"])

    for row_number, (_, values) in enumerate(source.iterrows(), start=2):
        for column_index, name in enumerate(columns, start=1):
            cell = sheet.cell(row=row_number, column=column_index, value=_excel_value(name, values[name]))
            cell.font = Font(name="Calibri", size=11)
            if "Time" in name and cell.value is not None:
                cell.number_format = "h:mm AM/PM"
            if name in {"Start Location", "End Location"}:
                cell.alignment = Alignment(wrap_text=True, vertical="center")

    widths = {
        "Asn#": 16,
        "Block": 12,
        "Route": 12,
        "Direction": 14,
        "Start Time": 14,
        "Start Location": 42,
        "End Location": 42,
        "End Time": 14,
        "Type": 14,
        "Garage": 14,
        "Veh1": 12,
        "Duty": 12,
    }
    for index, name in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(name, 16)
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 22
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
